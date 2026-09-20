import urllib.parse
import urllib.request
import psycopg2
import json
import re
import os
import threading
import time
from datetime import datetime, timezone, timedelta
import json as json_module
import base64
import csv
import xml.etree.ElementTree as ET
from io import BytesIO, StringIO

from dotenv import load_dotenv

load_dotenv()
from werkzeug.security import generate_password_hash, check_password_hash

from flask import (
    Flask,
    request,
    jsonify,
    send_from_directory,
    session,
    Response,
    stream_with_context
)

from google import genai
from google.genai import types

from chat_history import (
    add_chat,
    load_history,
    delete_chat,
    rename_chat,
    pin_chat,
    move_chat
)

from memory import (
    load_memory,
    update_memory
)

from bs4 import BeautifulSoup

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


# ==========================================
# DATA FOLDER
# ==========================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.join(BASE_DIR, "..", "data")
USERS_FILE = os.path.join(DATA_DIR, "users.json")

os.makedirs(DATA_DIR, exist_ok=True)


def load_users():

    if not os.path.exists(USERS_FILE):
        return []

    try:

        with open(
            USERS_FILE,
            "r",
            encoding="utf-8"
        ) as file:

            return json.load(file)

    except Exception:

        return []


def save_users(users):

    with open(
        USERS_FILE,
        "w",
        encoding="utf-8"
    ) as file:

        json.dump(
            users,
            file,
            indent=4,
            ensure_ascii=False
        )


# ==========================================
# APP
# ==========================================

SECRET_KEY = os.environ.get(
    "FLASK_SECRET_KEY"
)

if not SECRET_KEY:

    raise RuntimeError(
        "FLASK_SECRET_KEY is not set"
    )


app = Flask(__name__)

app.secret_key = SECRET_KEY
app.config["MAX_CONTENT_LENGTH"] = 14 * 1024 * 1024

app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = (
    os.environ.get("RENDER") == "true"
)


# ==========================================
# DATABASE
# ==========================================

def get_db_connection():

    database_url = os.environ.get(
        "DATABASE_URL"
    )

    if not database_url:

        raise RuntimeError(
            "DATABASE_URL is not set"
        )

    return psycopg2.connect(
        database_url
    )


def init_db():

    conn = get_db_connection()

    try:

        with conn.cursor() as cursor:

            cursor.execute("""
                CREATE TABLE IF NOT EXISTS memories (
                    username TEXT NOT NULL,
                    key TEXT NOT NULL,
                    value TEXT NOT NULL,
                    PRIMARY KEY (username, key)
                )
            """)

            cursor.execute("""
                CREATE TABLE IF NOT EXISTS users (
                    id SERIAL PRIMARY KEY,
                    username TEXT UNIQUE NOT NULL,
                    password TEXT NOT NULL
                )
            """)

            cursor.execute("""
                CREATE TABLE IF NOT EXISTS chat_history (
                    id SERIAL PRIMARY KEY,
                    username TEXT NOT NULL,
                    title TEXT NOT NULL,
                    messages JSONB NOT NULL,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)

        conn.commit()

    finally:

        conn.close()


# ==========================================
# GEMINI
# ==========================================

MODEL = "gemini-3.6-flash"

FALLBACK_MODEL = "gemini-3.5-flash-lite"

client = genai.Client()


# ==========================================
# MEMORY HELPERS
# ==========================================

def clean_value(value):

    value = str(value).strip()

    value = re.sub(
        r"\s+(hai|hain|is)$",
        "",
        value,
        flags=re.IGNORECASE
    )

    value = value.rstrip(
        ".,!?"
    ).strip()

    return value


def save_memory(
    key,
    value,
    user_memory
):

    username = session.get(
        "username"
    )

    if not username:
        return

    value = clean_value(
        value
    )

    if not value:
        return

    update_memory(
        key,
        value,
        username
    )

    user_memory[key] = value

    print(
        f"Memory saved: {key} = {value}"
    )


# ==========================================
# AUTOMATIC MEMORY DETECTION
# ==========================================


def detect_memory(message, user_memory):

    text = re.sub(r"\s+", " ", message.strip())
    lower = text.lower()

    # Normalize common Hinglish spellings so natural variations work.
    lower = lower.replace("favourite", "favorite")

    patterns = [

        # NAME
        (r"^(?:mera naam|my name is|my name's)\s+(.+?)(?:\s+hai)?[.!?]?$", "name"),
        (r"^(?:main|mai|i am|i'm)\s+([A-Za-z][A-Za-z\s]{1,40})[.!?]?$", "name"),
        (r"^(?:mujhe|you can)\s+(?:kunal|mujhe)\s+(?:bulao|call me)[.!?]?$", "name"),
        (r"^(?:call me|you can call me)\s+([A-Za-z][A-Za-z\s]{1,40})[.!?]?$", "name"),
        (r"^(?:mujhe)\s+(.+?)\s+(?:ke naam se|naam se)\s+bulao[.!?]?$", "name"),

        # COLLEGE
        (r"^(?:main|mai)\s+(.+?)\s+me\s+padh(?:ta|ti)\s+hoon[.!?]?$", "college"),
        (r"^(?:i study at|i study in|i go to)\s+(.+?)[.!?]?$", "college"),
        (r"^(?:mera|meri)\s+college\s+(.+?)(?:\s+hai)?[.!?]?$", "college"),

        # CITY
        (r"^(?:main|mai)\s+(.+?)\s+me\s+rehta\s+hoon[.!?]?$", "city"),
        (r"^(?:meri city|my city)\s+(.+?)(?:\s+hai)?[.!?]?$", "city"),
        (r"^(?:i live in|i am from|i'm from)\s+(.+?)[.!?]?$", "city"),

        # FAVORITE GAME
        (r"^(?:mera|meri)\s+favorite\s+game\s+(.+?)(?:\s+hai)?[.!?]?$", "favorite game"),
        (r"^my\s+favorite\s+game\s+is\s+(.+?)[.!?]?$", "favorite game"),
        (r"^(?:mujhe|i)\s+(.+?)\s+(?:game\s+)?pasand\s+hai[.!?]?$", "favorite game"),

        # FAVORITE COLOR
        (r"^(?:mera|meri)\s+favorite\s+color\s+(.+?)(?:\s+hai)?[.!?]?$", "favorite color"),
        (r"^my\s+favorite\s+color\s+is\s+(.+?)[.!?]?$", "favorite color"),

        # PROGRAMMING LANGUAGE
        (r"^(?:mera|meri)\s+favorite\s+programming language\s+(.+?)(?:\s+hai)?[.!?]?$", "favorite programming language"),
        (r"^my\s+favorite\s+programming language\s+is\s+(.+?)[.!?]?$", "favorite programming language"),
        (r"^(?:i like|i love)\s+(python|java|javascript|c\+\+|c|php|go|rust)[.!?]?$", "favorite programming language"),

        # GOAL
        (r"^mera goal\s+(.+?)(?:\s+hai)?[.!?]?$", "goal"),
        (r"^my goal is\s+(.+?)[.!?]?$", "goal"),
        (r"^(?:i want to become|i want to be|mera aim)\s+(.+?)[.!?]?$", "goal"),

        # GENERAL PREFERENCE
        (r"^mujhe\s+(.+?)\s+pasand\s+hai[.!?]?$", "preference"),
        (r"^i like\s+(.+?)[.!?]?$", "preference")
    ]

    for pattern, key in patterns:
        match = re.search(pattern, lower, re.IGNORECASE)
        if not match:
            continue

        value = clean_value(match.group(1))
        if not value:
            return None, None

        # Keep the user's original capitalization for stored values.
        original_match = re.search(pattern, text, re.IGNORECASE)
        if original_match:
            value = clean_value(original_match.group(1))

        # Guard against accidentally storing conversational filler.
        if len(value) > 120:
            return None, None

        if key == "preference":
            save_memory("favorite", value, user_memory)
        else:
            save_memory(key, value, user_memory)

        print(f"SMART MEMORY DETECTED: {key} = {value}")
        return None, None

    return None, None

# ==========================================
# PERSONAL QUESTIONS
# ==========================================

def personal_answer(
    message,
    user_memory
):

    text = message.lower().strip()

    text = re.sub(
        r"[?.!]",
        "",
        text
    )

    # AI NAME

    if any(
        x in text
        for x in [
            "tumhara naam kya hai",
            "tumhare naam kya hai",
            "aapka naam kya hai",
            "aapke naam kya hai"
        ]
    ):

        return (
            "Mera naam Normal Chat hai."
        )

    # USER NAME

    if any(
        x in text
        for x in [
            "mera naam kya hai",
            "my name kya hai",
            "what is my name"
        ]
    ):

        if "name" in user_memory:

            return (
                f"Tumhara naam "
                f"{user_memory['name']} hai."
            )

        return (
            "Mujhe abhi tumhara naam nahi pata."
        )

    # FAVORITE GAME

    if "mera favorite game kya hai" in text:

        if "favorite game" in user_memory:

            return (
                f"Tumhara favorite game "
                f"{user_memory['favorite game']} hai."
            )

        return (
            "Mujhe abhi tumhara favorite game nahi pata."
        )

    # FAVORITE COLOR

    if "mera favorite color kya hai" in text:

        if "favorite color" in user_memory:

            return (
                f"Tumhara favorite color "
                f"{user_memory['favorite color']} hai."
            )

        return (
            "Mujhe abhi tumhara favorite color nahi pata."
        )

    # FAVORITE

    if "mera favorite kya hai" in text:

        if "favorite" in user_memory:

            return (
                f"Tumhe "
                f"{user_memory['favorite']} pasand hai."
            )

        return (
            "Mujhe abhi tumhara favorite nahi pata."
        )

    # COLLEGE

    if any(
        x in text
        for x in [
            "mera college kya hai",
            "mera college ka kya naam hai",
            "mere college ka kya naam hai",
            "what is my college"
        ]
    ):

        if "college" in user_memory:

            return (
                f"Tumhara college "
                f"{user_memory['college']} hai."
            )

        return (
            "Mujhe abhi tumhara college nahi pata."
        )

    # GOAL

    if "mera goal kya hai" in text:

        if "goal" in user_memory:

            return (
                f"Tumhara goal "
                f"{user_memory['goal']} hai."
            )

        return (
            "Mujhe abhi tumhara goal nahi pata."
        )

    # CITY

    if "meri city kya hai" in text:

        if "city" in user_memory:

            return (
                f"Tumhari city "
                f"{user_memory['city']} hai."
            )

        return (
            "Mujhe abhi tumhari city nahi pata."
        )

    return None


# ==========================================
# SIGNUP
# ==========================================

@app.route(
    "/signup",
    methods=["POST"]
)
def signup():

    try:

        data = request.get_json(
            silent=True
        )

        if not data:

            return jsonify({
                "error": "Request data missing"
            }), 400

        username = str(
            data.get(
                "username",
                ""
            )
        ).strip()

        password = str(
            data.get(
                "password",
                ""
            )
        )

        if not username or not password:

            return jsonify({
                "error":
                "Username and password are required"
            }), 400

        if len(username) < 3:

            return jsonify({
                "error":
                "Username must be at least 3 characters"
            }), 400

        if len(password) < 8:

            return jsonify({
                "error":
                "Password must be at least 8 characters"
            }), 400

        password_hash = generate_password_hash(
            password
        )

        conn = get_db_connection()

        try:

            with conn.cursor() as cursor:

                cursor.execute(
                    """
                    INSERT INTO users
                    (username, password)
                    VALUES (%s, %s)
                    """,
                    (
                        username,
                        password_hash
                    )
                )

            conn.commit()

        except psycopg2.errors.UniqueViolation:

            conn.rollback()

            return jsonify({
                "error":
                "Username already exists"
            }), 409

        finally:

            conn.close()

        session["username"] = username

        return jsonify({
            "message":
            "Signup successful",
            "username":
            username
        }), 201

    except Exception as error:

        print(
            "Backend error:",
            repr(error)
        )

        return jsonify({
            "error":
            str(error)
        }), 500


# ==========================================
# LOGIN
# ==========================================

@app.route(
    "/login",
    methods=["POST"]
)
def login():

    try:

        data = request.get_json(
            silent=True
        )

        if not data:

            return jsonify({
                "error":
                "Request data missing"
            }), 400

        username = str(
            data.get(
                "username",
                ""
            )
        ).strip()

        password = str(
            data.get(
                "password",
                ""
            )
        )

        if not username or not password:

            return jsonify({
                "error":
                "Username and password are required"
            }), 400

        conn = get_db_connection()

        try:

            with conn.cursor() as cursor:

                cursor.execute(
                    """
                    SELECT username, password
                    FROM users
                    WHERE LOWER(username)
                    = LOWER(%s)
                    """,
                    (username,)
                )

                user = cursor.fetchone()

        finally:

            conn.close()

        if not user:

            return jsonify({
                "error":
                "Invalid username or password"
            }), 401

        stored_username, stored_password = user

        if not check_password_hash(
            stored_password,
            password
        ):

            return jsonify({
                "error":
                "Invalid username or password"
            }), 401

        session["username"] = stored_username

        return jsonify({
            "message":
            "Login successful",
            "username":
            stored_username
        }), 200

    except Exception as error:

        import traceback

        print(
            "========== LOGIN ERROR =========="
        )

        print(
            "ERROR:",
            repr(error)
        )

        traceback.print_exc()

        print(
            "================================="
        )

        return jsonify({
            "error":
            "Login error",
            "details":
            str(error)
        }), 500


# ==========================================
# LOGOUT
# ==========================================

@app.route(
    "/logout",
    methods=["POST"]
)
def logout():

    session.clear()

    return jsonify({
        "message":
        "Logged out successfully"
    })


# ==========================================
# CURRENT USER
# ==========================================

@app.route(
    "/me",
    methods=["GET"]
)
def me():

    username = session.get(
        "username"
    )

    if username:

        return jsonify({
            "logged_in":
            True,
            "username":
            username
        })

    return jsonify({
        "logged_in":
        False
    })


# ==========================================
# HOME
# ==========================================

@app.route("/")
def home():

    return send_from_directory(
        os.path.join(
            BASE_DIR,
            "frontend"
        ),
        "index.html"
    )


# ==========================================
# HEALTH CHECK
# ==========================================

@app.route("/health")
def health():

    return jsonify({
        "status":
        "ok"
    })


# ==========================================
# REAL-TIME WEATHER
# ==========================================

def get_weather(city):

    try:

        city_encoded = urllib.parse.quote(
            city
        )

        geo_url = (
            "https://geocoding-api.open-meteo.com/v1/search"
            f"?name={city_encoded}"
            "&count=1"
            "&language=en"
            "&format=json"
        )

        with urllib.request.urlopen(
            geo_url,
            timeout=10
        ) as response:

            geo_data = json.loads(
                response.read().decode(
                    "utf-8"
                )
            )

        if not geo_data.get(
            "results"
        ):

            return (
                f"Sorry, mujhe '{city}' "
                "naam ki city nahi mili."
            )

        location = geo_data[
            "results"
        ][0]

        latitude = location[
            "latitude"
        ]

        longitude = location[
            "longitude"
        ]

        city_name = location[
            "name"
        ]

        country = location.get(
            "country",
            ""
        )

        timezone = location.get(
            "timezone",
            "auto"
        )

        weather_url = (
            "https://api.open-meteo.com/v1/forecast"
            f"?latitude={latitude}"
            f"&longitude={longitude}"
            "&current="
            "temperature_2m,"
            "relative_humidity_2m,"
            "apparent_temperature,"
            "precipitation,"
            "weather_code,"
            "wind_speed_10m"
            f"&timezone="
            f"{urllib.parse.quote(timezone)}"
        )

        with urllib.request.urlopen(
            weather_url,
            timeout=10
        ) as response:

            weather_data = json.loads(
                response.read().decode(
                    "utf-8"
                )
            )

        current = weather_data.get(
            "current",
            {}
        )

        temperature = current.get(
            "temperature_2m"
        )

        feels_like = current.get(
            "apparent_temperature"
        )

        humidity = current.get(
            "relative_humidity_2m"
        )

        rain = current.get(
            "precipitation"
        )

        wind = current.get(
            "wind_speed_10m"
        )

        weather_code = current.get(
            "weather_code"
        )

        conditions = {

            0: "Clear sky",
            1: "Mainly clear",
            2: "Partly cloudy",
            3: "Overcast",
            45: "Foggy",
            48: "Foggy",
            51: "Light drizzle",
            53: "Drizzle",
            55: "Heavy drizzle",
            56: "Freezing drizzle",
            57: "Heavy freezing drizzle",
            61: "Light rain",
            63: "Rain",
            65: "Heavy rain",
            66: "Freezing rain",
            67: "Heavy freezing rain",
            71: "Light snow",
            73: "Snow",
            75: "Heavy snow",
            77: "Snow grains",
            80: "Light rain showers",
            81: "Rain showers",
            82: "Heavy rain showers",
            85: "Light snow showers",
            86: "Heavy snow showers",
            95: "Thunderstorm",
            96: "Thunderstorm with hail",
            99: "Heavy thunderstorm with hail"
        }

        condition = conditions.get(
            weather_code,
            "Unknown weather"
        )

        return (
            f"🌤️ Weather in "
            f"{city_name}, {country}\n\n"
            f"🌡️ Temperature: "
            f"{temperature}°C\n"
            f"🤒 Feels like: "
            f"{feels_like}°C\n"
            f"☁️ Condition: "
            f"{condition}\n"
            f"💧 Humidity: "
            f"{humidity}%\n"
            f"🌧️ Precipitation: "
            f"{rain} mm\n"
            f"💨 Wind speed: "
            f"{wind} km/h"
        )

    except Exception as error:

        print(
            "WEATHER ERROR:",
            repr(error)
        )

        return (
            "Sorry, weather service "
            "abhi available nahi hai."
        )


# ==========================================
# WEB SEARCH
# ==========================================

def _search_domain(url):

    try:
        hostname = urllib.parse.urlparse(str(url or "")).netloc.lower()
        if hostname.startswith("www."):
            hostname = hostname[4:]
        return hostname
    except Exception:
        return ""


def web_search(
    query,
    max_results=5
):

    """Search current web/news results through Google News RSS.

    Gemini is used only for answer generation; search results are fetched
    independently so a Gemini Search-grounding quota does not break search.
    """

    try:

        query = re.sub(r"\s+", " ", str(query).strip())

        if not query:
            return []

        # Add the current India date for time-sensitive queries so the
        # RSS search is anchored to today's news rather than older results.
        ist = timezone(timedelta(hours=5, minutes=30))
        current_date = datetime.now(ist).strftime("%d %B %Y")
        current_query = query
        current_markers = (
            "latest", "current", "today", "tonight", "recent",
            "aaj", "ajj", "abhi", "taaza", "taza", "filhaal",
            "news", "update", "live"
        )
        if any(marker in query.lower() for marker in current_markers):
            current_query = f"{query} {current_date}"

        search_url = (
            "https://news.google.com/rss/search?q="
            + urllib.parse.quote(current_query)
            + "&hl=en-IN&gl=IN&ceid=IN:en"
        )

        req = urllib.request.Request(
            search_url,
            headers={
                "User-Agent":
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0 Safari/537.36",
                "Accept-Language": "en-IN,en;q=0.9"
            }
        )

        with urllib.request.urlopen(
            req,
            timeout=15
        ) as response:

            xml_data = response.read()

        root = ET.fromstring(xml_data)

        results = []
        seen_urls = set()
        seen_titles = set()

        for item in root.findall(".//item"):

            title = (item.findtext("title") or "").strip()
            url = (item.findtext("link") or "").strip()
            description = (item.findtext("description") or "").strip()
            pub_date = (item.findtext("pubDate") or "").strip()

            source_element = item.find("source")
            source_name = (
                source_element.text.strip()
                if source_element is not None and source_element.text
                else ""
            )

            if not title or not url:
                continue

            normalized_url = url.rstrip("/").lower()
            normalized_title = re.sub(r"\s+", " ", title.lower())

            if normalized_url in seen_urls or normalized_title in seen_titles:
                continue

            seen_urls.add(normalized_url)
            seen_titles.add(normalized_title)

            # RSS descriptions can contain simple HTML markup.
            snippet = re.sub(r"<[^>]+>", " ", description)
            snippet = re.sub(r"\s+", " ", snippet).strip()

            results.append({
                "title": title[:300],
                "url": url,
                "domain": _search_domain(url),
                "snippet": snippet[:800],
                "source": source_name[:200],
                "published": pub_date[:100]
            })

            if len(results) >= max_results:
                break

        print(
            "WEB SEARCH:",
            query,
            "RESULTS:",
            len(results)
        )

        for index, result in enumerate(results, start=1):
            print(
                f"  SOURCE {index}: "
                f"{result['domain']} | "
                f"{result['title']}"
            )

        return results

    except Exception as error:

        print(
            "WEB SEARCH ERROR:",
            repr(error)
        )

        return []


def format_web_results(results):

    if not results:
        return ""

    lines = []

    for index, result in enumerate(
        results,
        start=1
    ):

        source_line = result.get("source", "")
        published_line = result.get("published", "")

        lines.append(
            f"{index}. {result['title']}\n"
            f"Source: {source_line or result.get('domain', '')}\n"
            f"Published: {published_line}\n"
            f"URL: {result['url']}\n"
            f"{result['snippet']}"
        )

    return "\n\n".join(lines)


def build_smart_search_query(message, history):
    """Build a cleaner search query using the current message and recent context."""
    current = re.sub(r"\s+", " ", str(message).strip())

    current = re.sub(
        r"^(please\s+|can you\s+|could you\s+|tell me\s+|mujhe\s+|zara\s+|batao\s+)",
        "",
        current,
        flags=re.IGNORECASE
    ).strip()

    context_parts = []
    if isinstance(history, list):
        for msg in history[-8:]:
            if not isinstance(msg, dict):
                continue
            role = str(msg.get("role", "")).lower()
            content = str(msg.get("content", "")).strip()
            if role == "user" and content:
                context_parts.append(content)

    followup = bool(re.match(
        r"^(and|also|what about|how about|aur|iske|ispe|uska|uske|ye|yeh|that|this|same|why|when|where|who|how)\b",
        current,
        re.IGNORECASE
    ))

    if followup and context_parts:
        query = f"{context_parts[-1]} {current}"
    else:
        query = current

    query = re.sub(r"\s+", " ", query).strip()
    return query[:500]


# ==========================================

def generate_ai_reply(
    prompt,
    use_web_search=False,
    image_data=None
):

    last_error = None

    # Web search is fetched independently and already included in the prompt.
    # Do not enable Gemini Search grounding here.
    models_to_try = [
        MODEL,
        FALLBACK_MODEL
    ]

    for current_model in models_to_try:

        for attempt in range(2):

            try:

                print(
                    f"Trying Gemini model: {current_model} "
                    f"(attempt {attempt + 1})"
                )

                if image_data:
                    image_part = types.Part.from_bytes(
                        data=image_data["bytes"],
                        mime_type=image_data["mime_type"]
                    )

                    contents = [
                        prompt,
                        image_part
                    ]

                    print(
                        "GEMINI IMAGE INPUT: READY",
                        image_data["mime_type"],
                        len(image_data["bytes"]),
                        "bytes"
                    )
                else:
                    contents = prompt

                response = client.models.generate_content(
                    model=current_model,
                    contents=contents
                )

                if response and response.text:

                    if use_web_search:
                        print(
                            "WEB RESULTS SENT TO GEMINI: SUCCESS"
                        )

                    return response.text.strip()

                raise Exception(
                    "Gemini ne empty response diya."
                )

            except Exception as error:

                last_error = error
                error_text = str(error)

                print(
                    f"Gemini error on {current_model}: {error_text}"
                )

                if (
                    "503" in error_text
                    or "UNAVAILABLE" in error_text
                    or "429" in error_text
                ):
                    time.sleep(2)
                    continue

                break

    if last_error:
        raise last_error

    raise Exception(
        "Gemini ne response nahi diya."
    )


# ==========================================

def is_image_edit_request(message):
    """Detect requests that ask to modify an attached image."""
    text = re.sub(r"\s+", " ", str(message or "").strip().lower())
    if not text:
        return False

    edit_patterns = [
        r"\b(change|modify|edit|alter|replace|remove|add|make|turn|convert|transform)\b.*\b(background|color|colour|object|person|sky|hair|dress|shirt|shirt|wall)\b",
        r"\b(background|bg)\s+(?:ko|to|into|mein|me)\b",
        r"\b(remove|delete|erase)\b.*\b(from|image|photo|picture)\b",
        r"\b(add|put|insert)\b.*\b(to|in|on|the image|the photo|the picture)\b",
        r"\bmake\s+the\s+background\b",
        r"\bbackground\s+(?:blue|red|green|black|white|yellow|pink|purple|orange|grey|gray)\b",
        r"\b(?:blue|red|green|black|white|yellow|pink|purple|orange|grey|gray)\s+background\b",
        r"\bchange\s+.*\bcolor\b",
        r"\bbackground\s+color\b",
        r"\bbackground\s+colour\b",
    ]

    return any(re.search(pattern, text, re.IGNORECASE) for pattern in edit_patterns)


def generate_edited_image(image_data, edit_prompt):
    """Edit an attached image using Gemini 3.1 Flash Image."""
    if not image_data or not image_data.get("bytes"):
        raise ValueError("Image attachment is required for image editing.")

    encoded_image = base64.b64encode(image_data["bytes"]).decode("utf-8")

    prompt = (
        "Edit the provided image according to the user's request. "
        "Preserve the main subject, composition, proportions, and important details "
        "unless the user explicitly asks to change them. Make only the requested edit "
        "and keep the result natural and high quality.\n\n"
        f"User request: {edit_prompt}"
    )

    print("IMAGE EDIT REQUEST:", edit_prompt)
    print("IMAGE EDIT MODEL: gemini-3.1-flash-image")

    interaction = client.interactions.create(
        model="gemini-3.1-flash-image",
        input=[
            {"type": "text", "text": prompt},
            {
                "type": "image",
                "data": encoded_image,
                "mime_type": image_data["mime_type"],
            },
        ],
        response_format={
            "type": "image",
            "mime_type": "image/jpeg",
        },
    )

    output_image = getattr(interaction, "output_image", None)
    if output_image is None or not getattr(output_image, "data", None):
        raise Exception("Gemini image model ne edited image return nahi ki.")

    output_data = output_image.data
    if isinstance(output_data, bytes):
        output_data = base64.b64encode(output_data).decode("utf-8")

    return {
        "data": output_data,
        "mime_type": getattr(output_image, "mime_type", None) or "image/png",
    }


# ==========================================
# FILE / DOCUMENT EXTRACTION
# ==========================================

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_DOCUMENT_CHARS = 60000

SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".pdf",
    ".docx",
    ".xlsx",
}


def _truncate_document_text(text):

    text = str(text or "").strip()

    if len(text) <= MAX_DOCUMENT_CHARS:
        return text

    return (
        text[:MAX_DOCUMENT_CHARS]
        + "\n\n[Document text truncated for processing.]"
    )


def extract_uploaded_document(file_name, mime_type, file_bytes):
    """Extract useful text from supported document formats."""

    file_name = str(file_name or "file").strip()
    mime_type = str(mime_type or "").strip().lower()
    extension = os.path.splitext(file_name)[1].lower()

    if not file_bytes:
        raise ValueError("Uploaded file is empty.")

    if len(file_bytes) > MAX_UPLOAD_BYTES:
        raise ValueError("File must be smaller than 10 MB.")

    # Plain text / Markdown / CSV
    if (
        extension in {".txt", ".md", ".csv"}
        or mime_type.startswith("text/")
    ):
        try:
            text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = file_bytes.decode("utf-8", errors="replace")

        if extension == ".csv" or "csv" in mime_type:
            try:
                rows = list(csv.reader(StringIO(text)))
                lines = ["\t".join(row) for row in rows]
                text = "\n".join(lines)
            except Exception:
                pass

        return _truncate_document_text(text)

    # PDF
    if extension == ".pdf" or mime_type == "application/pdf":
        if PdfReader is None:
            raise RuntimeError(
                "PDF support is not installed. Run: pip install pypdf"
            )

        reader = PdfReader(BytesIO(file_bytes))
        pages = []

        for page_number, page in enumerate(reader.pages, start=1):
            try:
                page_text = page.extract_text() or ""
            except Exception as error:
                print(
                    f"PDF PAGE {page_number} EXTRACTION ERROR:",
                    repr(error)
                )
                page_text = ""

            if page_text.strip():
                pages.append(
                    f"[Page {page_number}]\n{page_text.strip()}"
                )

            current = "\n\n".join(pages)
            if len(current) >= MAX_DOCUMENT_CHARS:
                break

        return _truncate_document_text(
            "\n\n".join(pages)
        )

    # DOCX
    if (
        extension == ".docx"
        or mime_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        if Document is None:
            raise RuntimeError(
                "DOCX support is not installed. Run: pip install python-docx"
            )

        document = Document(BytesIO(file_bytes))
        parts = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                parts.append(text)

        # Include table contents too.
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells]
                parts.append("\t".join(values))

        return _truncate_document_text(
            "\n".join(parts)
        )

    # XLSX
    if extension == ".xlsx":
        if load_workbook is None:
            raise RuntimeError(
                "XLSX support is not installed. Run: pip install openpyxl"
            )

        workbook = load_workbook(
            filename=BytesIO(file_bytes),
            read_only=True,
            data_only=True
        )

        try:
            parts = []

            for worksheet in workbook.worksheets:
                parts.append(
                    f"[Sheet: {worksheet.title}]"
                )

                for row in worksheet.iter_rows(values_only=True):
                    values = [
                        "" if value is None else str(value)
                        for value in row
                    ]

                    if any(value.strip() for value in values):
                        parts.append("\t".join(values))

                    if len("\n".join(parts)) >= MAX_DOCUMENT_CHARS:
                        break

                if len("\n".join(parts)) >= MAX_DOCUMENT_CHARS:
                    break

            return _truncate_document_text(
                "\n".join(parts)
            )

        finally:
            workbook.close()

    raise ValueError(
        "Unsupported file type. Supported files: PDF, TXT, MD, CSV, DOCX, XLSX and images."
    )


# ==========================================
# CHAT
# ==========================================

@app.route(
    "/chat",
    methods=["POST"]
)
def chat():

    username = session.get(
        "username"
    )

    if not username:

        return jsonify({
            "error":
            "Login required"
        }), 401

    user_memory = load_memory(
        username
    )

    try:

        data = request.get_json(
            silent=True
        )

        if not data or "message" not in data:

            return jsonify({
                "error":
                "Message is missing"
            }), 400

        message = str(
            data["message"]
        ).strip()

        if not message:

            return jsonify({
                "error":
                "Message is empty"
            }), 400

        if len(message) > 5000:

            return jsonify({
                "error":
                "Message too long. "
                "Maximum 5000 characters allowed."
            }), 400

        response_style = str(
            data.get("responseStyle", "balanced")
        ).strip().lower()

        if response_style not in (
            "concise",
            "balanced",
            "detailed"
        ):
            response_style = "balanced"

        # ======================================
        # FILE / IMAGE ATTACHMENT
        # ======================================

        image_data = None
        document_text = ""
        uploaded_file_name = ""
        uploaded_file_mime = ""

        image_payload = data.get("image")
        file_payload = data.get("file")

        attachment_payload = (
            image_payload
            if image_payload
            else file_payload
        )

        if attachment_payload:

            if not isinstance(attachment_payload, dict):
                return jsonify({
                    "error": "Invalid file attachment"
                }), 400

            import base64

            uploaded_file_name = str(
                attachment_payload.get(
                    "name",
                    "file"
                )
            ).strip()

            uploaded_file_mime = str(
                attachment_payload.get(
                    "mimeType",
                    ""
                )
            ).strip().lower()

            data_url = str(
                attachment_payload.get(
                    "dataUrl",
                    ""
                )
            ).strip()

            if (
                not data_url.startswith("data:")
                or ";base64," not in data_url
            ):
                return jsonify({
                    "error": "Invalid file data."
                }), 400

            try:
                encoded = data_url.split(
                    ";base64,",
                    1
                )[1]

                file_bytes = base64.b64decode(
                    encoded,
                    validate=True
                )

            except Exception:
                return jsonify({
                    "error": "File data could not be read."
                }), 400

            if not file_bytes:
                return jsonify({
                    "error": "Uploaded file is empty."
                }), 400

            if len(file_bytes) > MAX_UPLOAD_BYTES:
                return jsonify({
                    "error": "File must be smaller than 10 MB."
                }), 400

            print(
                "FILE ATTACHMENT RECEIVED:",
                uploaded_file_name,
                uploaded_file_mime,
                len(file_bytes),
                "bytes"
            )

            if uploaded_file_mime.startswith("image/"):

                if len(file_bytes) > 8 * 1024 * 1024:
                    return jsonify({
                        "error": "Image must be smaller than 8 MB."
                    }), 400

                image_data = {
                    "bytes": file_bytes,
                    "mime_type": uploaded_file_mime
                }

            else:

                try:
                    document_text = extract_uploaded_document(
                        uploaded_file_name,
                        uploaded_file_mime,
                        file_bytes
                    )
                except Exception as error:
                    print(
                        "DOCUMENT EXTRACTION ERROR:",
                        repr(error)
                    )
                    return jsonify({
                        "error": str(error)
                    }), 400

                if not document_text.strip():
                    return jsonify({
                        "error": "Document mein readable text nahi mila."
                    }), 400

        # ======================================
        # IMAGE EDITING
        # ======================================

        if image_data and is_image_edit_request(message):
            edited_image = generate_edited_image(
                image_data,
                message
            )

            def image_edit_stream():
                yield (
                    "data: "
                    + json_module.dumps({
                        "type": "image",
                        "mimeType": edited_image["mime_type"],
                        "data": edited_image["data"],
                    })
                    + "\n\n"
                )

                yield (
                    "data: "
                    + json_module.dumps({
                        "type": "chunk",
                        "text": "Image edit complete."
                    })
                    + "\n\n"
                )

                yield (
                    "data: "
                    + json_module.dumps({
                        "type": "done"
                    })
                    + "\n\n"
                )

            return Response(
                stream_with_context(image_edit_stream()),
                mimetype="text/event-stream",
                headers={
                    "Cache-Control": "no-cache",
                    "X-Accel-Buffering": "no",
                    "Connection": "keep-alive"
                }
            )

        history_from_frontend = data.get(
            "history",
            []
        )

        # ======================================
        # WEB SEARCH DETECTION
        # ======================================

        web_search_needed = False

        message_lower = message.lower()

        search_triggers = [

            # Time-sensitive
            "latest",
            "current",
            "today",
            "tonight",
            "yesterday",
            "tomorrow",
            "recent",
            "recently",
            "right now",
            "as of now",
            "at the moment",

            # News
            "news",
            "breaking news",
            "what happened",
            "what's happening",
            "what is happening",
            "happening now",

            # Current status
            "who is the current",
            "who is currently",
            "current pm",
            "current president",
            "current chief minister",
            "current minister",
            "current ceo",

            # Prices
            "current price",
            "latest price",
            "price today",
            "stock price",
            "share price",
            "bitcoin price",
            "crypto price",

            # Updates
            "latest update",
            "latest updates",
            "current update",
            "new update",
            "recent update",
            "recent updates",

            # Time periods
            "this week",
            "this month",
            "this year",
            "this weekend",

            # Live
            "live score",
            "live news",
            "live update",
            "live updates"
        ]

        hindi_search_patterns = [

            r"\babhi\s+(?:kya|ka|ki|ke|chal)\b",
            r"\babhi\s+ka\s+news\b",
            r"\baaj\s+(?:ka|ki|ke)\b",
            r"\baaj\s+ki\s+news\b",
            r"\baaj\s+ka\s+news\b",
            r"\btaza\s+(?:khabar|news)\b",
            r"\btaaza\s+(?:khabar|news)\b",
            r"\bhaal\s+hi\s+mein\b",
            r"\bfilhaal\b",
            r"\bcurrently\b",
            r"\biss\s+samay\b",
            r"\bis\s+samay\b",
            r"\babhi\s+ka\s+update\b",
            r"\blatest\s+update\b"
        ]

        for pattern in hindi_search_patterns:

            if re.search(
                pattern,
                message_lower,
                re.IGNORECASE
            ):

                web_search_needed = True

                break

        if not web_search_needed:

            for trigger in search_triggers:

                if trigger in message_lower:

                    web_search_needed = True

                    break

        web_context = ""

        if web_search_needed:

            smart_query = build_smart_search_query(
                message,
                history_from_frontend
            )

            print(
                "WEB SEARCH REQUEST DETECTED:",
                message
            )

            print(
                "SMART SEARCH QUERY:",
                smart_query
            )

            search_results = web_search(
                smart_query,
                max_results=5
            )

            if not search_results and smart_query != message:
                print(
                    "SMART SEARCH EMPTY - RETRYING EXACT QUERY"
                )
                search_results = web_search(
                    message,
                    max_results=5
                )

            web_context = format_web_results(
                search_results
            )

            if web_context:
                print(
                    "WEB SEARCH CONTEXT READY:"
                    , len(search_results),
                    "sources"
                )
            else:
                print(
                    "WEB SEARCH CONTEXT EMPTY"
                )

        # ======================================
        # WEATHER DETECTION
        # ======================================

        weather_city = None

        weather_patterns = [

            r"^(.+?)\s+(?:ka|ki|ke)\s+"
            r"(?:weather|mausam)"
            r"(?:\s+(?:batao|bata|kya hai|"
            r"kaisa hai|kaisi hai))?[?.!]*$",

            r"^(.+?)\s+(?:mein|me)\s+"
            r"(?:weather|mausam)"
            r"(?:\s+(?:batao|bata|kaisa hai|"
            r"kaisi hai|kya hai))?[?.!]*$",

            r"^(?:weather|mausam)\s+"
            r"(?:in|of)\s+(.+?)[?.!]*$",

            r"^(?:what(?:'s| is)?\s+)?"
            r"(?:the\s+)?weather\s+"
            r"(?:in|of)\s+(.+?)[?.!]*$",

            r"^(?:temperature|temp)\s+"
            r"(?:in|of|ka|ki|ke)?\s*"
            r"(.+?)[?.!]*$",

            r"^(?:what is\s+)?"
            r"(?:the\s+)?temperature\s+"
            r"(?:in|of)\s+(.+?)[?.!]*$"
        ]

        for pattern in weather_patterns:

            match = re.search(
                pattern,
                message.strip(),
                re.IGNORECASE
            )

            if match:

                weather_city = match.group(
                    1
                ).strip(
                    " .?!,:'\""
                )

                break

        if weather_city:

            print(
                "WEATHER REQUEST DETECTED:",
                weather_city
            )

            weather_reply = get_weather(
                weather_city
            )

            def weather_stream():

                words = weather_reply.split(
                    " "
                )

                for i, word in enumerate(
                    words
                ):

                    chunk = word

                    if i < len(words) - 1:

                        chunk += " "

                    yield (
                        "data: "
                        +
                        json_module.dumps({
                            "type":
                            "chunk",
                            "text":
                            chunk
                        })
                        +
                        "\n\n"
                    )

                    time.sleep(
                        0.015
                    )

                yield (
                    "data: "
                    +
                    json_module.dumps({
                        "type":
                        "done"
                    })
                    +
                    "\n\n"
                )

            return Response(
                stream_with_context(
                    weather_stream()
                ),
                mimetype="text/event-stream",
                headers={
                    "Cache-Control":
                    "no-cache",
                    "X-Accel-Buffering":
                    "no",
                    "Connection":
                    "keep-alive"
                }
            )

        # ======================================
        # PERSONAL QUESTION
        # ======================================

        direct_reply = personal_answer(
            message,
            user_memory
        )

        if direct_reply:

            def direct_stream():

                words = direct_reply.split(
                    " "
                )

                for i, word in enumerate(
                    words
                ):

                    chunk = word

                    if i < len(words) - 1:

                        chunk += " "

                    yield (
                        "data: "
                        +
                        json_module.dumps({
                            "type":
                            "chunk",
                            "text":
                            chunk
                        })
                        +
                        "\n\n"
                    )

                    time.sleep(
                        0.02
                    )

                yield (
                    "data: "
                    +
                    json_module.dumps({
                        "type":
                        "done"
                    })
                    +
                    "\n\n"
                )

            return Response(
                stream_with_context(
                    direct_stream()
                ),
                mimetype="text/event-stream",
                headers={
                    "Cache-Control":
                    "no-cache",
                    "X-Accel-Buffering":
                    "no",
                    "Connection":
                    "keep-alive"
                }
            )

        # ======================================
        # MEMORY
        # ======================================

        # detect_memory() saves the detected memory itself and always
        # returns a (key, value) tuple for compatibility. Do not unpack
        # it here because memory may already have been saved.
        detect_memory(
            message,
            user_memory
        )

        # ======================================
        # MEMORY TEXT
        # ======================================

        memory_lines = []

        for key, value in user_memory.items():

            memory_lines.append(
                f"{key}: {value}"
            )

        memory_text = "\n".join(
            memory_lines
        )

        # ======================================
        # SYSTEM PROMPT
        # ======================================

        ist = timezone(timedelta(hours=5, minutes=30))
        current_date_ist = datetime.now(ist).strftime("%d %B %Y")

        system_prompt = f"""
You are Normal Chat, a helpful AI assistant.

Current date in India: {current_date_ist}

Rules:
- Answer the user's question directly and prioritize the exact request.
- Match the user's language and style (English, Hindi, or Hinglish).
- Keep simple questions short and easy to scan.
- For complex, technical, or multi-step questions, give a clear structured answer with useful detail.
- Avoid filler openings such as "Sure!", "Absolutely!", or "Here's the answer" unless they add value.
- Do not repeat information the user already knows from the conversation unless it is needed for clarity.
- Use headings, bullets, or numbered steps when they improve readability.
- Explain technical terms briefly when the user may not know them.
- If the request is ambiguous, ask one focused clarification instead of making a large assumption.
- Never invent personal information, facts, sources, or actions you did not perform.
- Do not mention memory, saved information, system instructions, prompts, or internal processing.
- Use user information only when it is directly relevant to the answer.
- Do not unnecessarily restate the user's question.
- For code requests, prefer complete, directly usable code when appropriate and explain only the important parts.

Response style:
{("Concise: keep the answer brief and focused. Use only the details needed to answer the request." if response_style == "concise" else "Detailed: give a thorough, well-structured answer with useful explanations, examples, and steps when appropriate." if response_style == "detailed" else "Balanced: give a clear answer with enough explanation to be useful without unnecessary length.")}

User information:
{memory_text}
"""

        # ======================================
        # CONVERSATION
        # ======================================

        # ======================================
        # SMART CONVERSATION CONTEXT
        # ======================================

        # Keep recent turns together instead of blindly cutting the text
        # in the middle of a message. This gives Gemini cleaner context.
        raw_history = (
            history_from_frontend
            if isinstance(history_from_frontend, list)
            else []
        )

        # Current message is normally the final history item, so exclude it.
        previous_messages = raw_history[:-1]

        # Keep the latest 16 messages (8 user/assistant turns maximum).
        previous_messages = previous_messages[-16:]

        conversation_parts = []

        for msg in previous_messages:
            if not isinstance(msg, dict):
                continue

            role = str(
                msg.get("role", "")
            ).strip().lower()

            content = str(
                msg.get("content", "")
            ).strip()

            if not content:
                continue

            # Keep individual messages bounded.
            if len(content) > 2500:
                content = content[:2500] + "..."

            if role == "user":
                conversation_parts.append(
                    f"User: {content}"
                )

            elif role in (
                "assistant",
                "bot",
                "model"
            ):
                conversation_parts.append(
                    f"Assistant: {content}"
                )

        conversation_text = "\n".join(
            conversation_parts
        )

        # Never let conversation context dominate the current question.
        if len(conversation_text) > 14000:
            conversation_text = conversation_text[-14000:]

        # Extra instructions improve follow-up understanding.
        # The model should resolve words such as 'it', 'that', 'this',
        # 'above', and 'same one' from the most recent relevant context.
        conversation_rules = ""

        if conversation_text:
            conversation_rules = """
Conversation continuity rules:
- Treat the recent conversation as active context, not as a new question.
- If the user asks a follow-up such as 'why?', 'how?', 'what about it?',
  'that one', 'same thing', or 'explain more', resolve the reference from
  the most recent relevant user/assistant messages.
- Do not ask the user to repeat information that is already clear from the
  recent conversation.
- If the reference is genuinely ambiguous, ask one short clarification
  question instead of guessing.
- The newest user message always has priority over older context.
"""

        # ======================================
        # GEMINI PROMPT
        # ======================================

        prompt = f"""
{system_prompt}

If web search results are provided below:

- Use them as the primary source for current, recent, or time-sensitive facts.
- For requests containing "today", "aaj", "latest", "current", or similar wording, do not answer from model memory.
- Use the publication dates in the search results and never replace a current date with an older date.
- The current India date is {current_date_ist}.
- If the search results are unavailable or clearly stale, say that live search did not return a reliable current result instead of inventing one.
- Do not invent facts.
- Do not dump raw search results.
- Answer naturally and directly.

Web Search Results:
{web_context}

{conversation_rules}
Recent Conversation:
{conversation_text}

Current User Message:
{message}

Image Attachment:
{("An image is attached. Inspect the image and answer the user's request using it." if image_data else "No image is attached.")}

Document Attachment:
{(f"A document named {uploaded_file_name} is attached. Use the extracted document text below to answer the user's request.\n\n{document_text}" if document_text else "No document is attached.")}
"""

        # ======================================
        # GENERATE RESPONSE
        # ======================================

        reply = generate_ai_reply(
            prompt,
            use_web_search=web_search_needed,
            image_data=image_data
        )

        if not reply:

            raise Exception(
                "AI ne empty response diya."
            )

        # ======================================
        # STREAM RESPONSE
        # ======================================

        def generate_stream():

            words = reply.split(
                " "
            )

            for i, word in enumerate(
                words
            ):

                chunk = word

                if i < len(words) - 1:

                    chunk += " "

                yield (
                    "data: "
                    +
                    json_module.dumps({
                        "type":
                        "chunk",
                        "text":
                        chunk
                    })
                    +
                    "\n\n"
                )

                time.sleep(
                    0.015
                )

            yield (
                "data: "
                +
                json_module.dumps({
                    "type":
                    "done"
                })
                +
                "\n\n"
            )

        return Response(
            stream_with_context(
                generate_stream()
            ),
            mimetype="text/event-stream",
            headers={
                "Cache-Control":
                "no-cache",
                "X-Accel-Buffering":
                "no",
                "Connection":
                "keep-alive"
            }
        )

    except Exception as chat_error:

        import traceback

        print(
            "========== CHAT ERROR =========="
        )

        print(
            "ERROR:",
            repr(chat_error)
        )

        traceback.print_exc()

        print(
            "================================"
        )

        return jsonify({
            "error":
            "Internal server error",
            "details":
            str(chat_error)
        }), 500


# ==========================================
# HISTORY
# ==========================================

@app.route(
    "/history",
    methods=["GET"]
)
def history():

    username = session.get(
        "username"
    )

    if not username:

        return jsonify({
            "error":
            "Login required"
        }), 401

    try:

        return jsonify({
            "history":
            load_history(username)
        })

    except Exception as error:

        import traceback

        print(
            "========== HISTORY ERROR =========="
        )

        print(
            "ERROR:",
            repr(error)
        )

        traceback.print_exc()

        print(
            "==================================="
        )

        return jsonify({
            "error":
            str(error)
        }), 500


# ==========================================
# NEW CHAT
# ==========================================

@app.route(
    "/new-chat",
    methods=["POST"]
)
def new_chat():

    username = session.get(
        "username"
    )

    if not username:

        return jsonify({
            "error":
            "Login required"
        }), 401

    try:

        data = request.get_json(
            silent=True
        ) or {}

        history_data = data.get(
            "history",
            []
        )

        if history_data:

            add_chat(
                history_data,
                username
            )

        return jsonify({
            "message":
            "New chat started"
        })

    except Exception as error:

        print(
            "New chat error:",
            repr(error)
        )

        return jsonify({
            "error":
            "Internal server error",
            "details":
            str(error)
        }), 500


# ==========================================
# BRANCH CONVERSATION
# ==========================================

@app.route(
    "/branch",
    methods=["POST"]
)
def branch_conversation():

    username = session.get(
        "username"
    )

    if not username:

        return jsonify({
            "error":
            "Login required"
        }), 401

    try:

        data = request.get_json(
            silent=True
        ) or {}

        history_data = data.get(
            "history",
            []
        )

        if not isinstance(
            history_data,
            list
        ) or not history_data:

            return jsonify({
                "error":
                "Conversation history is required"
            }), 400

        # Keep the branch bounded and only retain normal
        # user/assistant messages.
        clean_history = []

        for item in history_data[-32:]:

            if not isinstance(item, dict):
                continue

            role = str(
                item.get("role", "")
            ).strip().lower()

            content = str(
                item.get("content", "")
            ).strip()

            if role not in (
                "user",
                "assistant"
            ) or not content:
                continue

            clean_history.append({
                "role": role,
                "content": content[:5000]
            })

        if not clean_history:

            return jsonify({
                "error":
                "No valid messages to branch"
            }), 400

        add_chat(
            clean_history,
            username
        )

        return jsonify({
            "message":
            "Conversation branch created",
            "history":
            clean_history
        })

    except Exception as error:

        import traceback

        print(
            "========== BRANCH ERROR =========="
        )

        print(
            "ERROR:",
            repr(error)
        )

        traceback.print_exc()

        print(
            "==================================="
        )

        return jsonify({
            "error":
            "Internal server error",
            "details":
            str(error)
        }), 500


# ==========================================
# DELETE HISTORY
# ==========================================

@app.route(
    "/history/<int:index>",
    methods=["DELETE"]
)
def delete_history(index):

    username = session.get(
        "username"
    )

    if not username:

        return jsonify({
            "error":
            "Login required"
        }), 401

    try:

        success = delete_chat(
            index,
            username
        )

        if not success:

            return jsonify({
                "error":
                "Chat not found"
            }), 404

        return jsonify({
            "message":
            "Chat deleted"
        })

    except Exception as error:

        print(
            "Delete history error:",
            repr(error)
        )

        return jsonify({
            "error":
            "Internal server error"
        }), 500


# ==========================================
# RENAME HISTORY
# ==========================================

@app.route(
    "/history/<int:index>/rename",
    methods=["POST", "PUT"]
)
def rename_history(index):

    username = session.get(
        "username"
    )

    if not username:

        return jsonify({
            "error":
            "Login required"
        }), 401

    try:

        data = request.get_json(
            silent=True
        )

        if not data:

            return jsonify({
                "error":
                "Request data missing"
            }), 400

        title = data.get(
            "name"
        )

        if title is None:

            title = data.get(
                "title"
            )

        if title is None:

            return jsonify({
                "error":
                "Name is missing"
            }), 400

        title = str(
            title
        ).strip()

        if not title:

            return jsonify({
                "error":
                "Name cannot be empty"
            }), 400

        title = title[:40]

        success = rename_chat(
            index,
            title,
            username
        )

        if not success:

            return jsonify({
                "error":
                "Chat not found"
            }), 404

        return jsonify({
            "message":
            "Chat renamed successfully",
            "title":
            title
        })

    except Exception as error:

        print(
            "Rename error:",
            repr(error)
        )

        return jsonify({
            "error":
            "Internal server error"
        }), 500
# ==========================================
# PIN / UNPIN HISTORY
# ==========================================

@app.route(
    "/history/<int:index>/pin",
    methods=["POST", "PUT"]
)
def pin_history(index):

    username = session.get(
        "username"
    )

    if not username:

        return jsonify({
            "error":
            "Login required"
        }), 401

    try:

        data = request.get_json(
            silent=True
        ) or {}

        pinned = data.get(
            "pinned",
            True
        )

        if isinstance(
            pinned,
            str
        ):

            pinned = (
                pinned.lower()
                in ("true", "1", "yes", "on")
            )

        success = pin_chat(
            index,
            bool(pinned),
            username
        )

        if not success:

            return jsonify({
                "error":
                "Chat not found"
            }), 404

        return jsonify({
            "message":
            (
                "Chat pinned"
                if pinned
                else
                "Chat unpinned"
            ),
            "pinned":
            bool(pinned)
        })

    except Exception as error:

        print(
            "Pin history error:",
            repr(error)
        )

        return jsonify({
            "error":
            "Internal server error"
        }), 500
# ==========================================
# MOVE CHAT TO FOLDER
# ==========================================

@app.route(
    "/history/<int:index>/folder",
    methods=["POST", "PUT"]
)
def move_history_folder(index):

    username = session.get(
        "username"
    )

    if not username:

        return jsonify({
            "error":
            "Login required"
        }), 401

    try:

        data = request.get_json(
            silent=True
        ) or {}

        folder = data.get(
            "folder",
            "General"
        )

        folder = str(
            folder
        ).strip()

        if not folder:

            folder = "General"

        folder = folder[:40]

        success = move_chat(
            index,
            folder,
            username
        )

        if not success:

            return jsonify({
                "error":
                "Chat not found"
            }), 404

        return jsonify({
            "message":
            "Chat moved successfully",
            "folder":
            folder
        })

    except Exception as error:

        print(
            "Move folder error:",
            repr(error)
        )

        return jsonify({
            "error":
            "Internal server error"
        }), 500
# ==========================================
# DATABASE STARTUP
# ==========================================

_db_initialized = False

_db_init_lock = threading.Lock()


@app.before_request
def ensure_database():

    global _db_initialized

    if (
        request.path == "/"
        or
        request.path.startswith(
            "/static/"
        )
    ):

        return None

    if _db_initialized:

        return None

    with _db_init_lock:

        if _db_initialized:

            return None

        try:

            init_db()

            _db_initialized = True

            print(
                "PostgreSQL database initialized successfully."
            )

        except Exception as error:

            print(
                "========== DATABASE INIT ERROR =========="
            )

            print(
                "ERROR:",
                repr(error)
            )

            print(
                "========================================="
            )

            return jsonify({
                "error":
                "Database connection failed",
                "details":
                str(error)
            }), 500

    return None


# ==========================================
# START
# ==========================================

if __name__ == "__main__":

    port = int(
        os.environ.get(
            "PORT",
            5000
        )
    )

    app.run(
        host="0.0.0.0",
        port=port,
        debug=False,
        threaded=True
    )